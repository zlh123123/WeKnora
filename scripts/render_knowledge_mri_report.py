#!/usr/bin/env python3
"""Render the submission Markdown to a Chinese PDF and editable DOCX.
Requires reportlab and python-docx; --font accepts an installed TTF/TTC.
"""
import argparse
import html
import re
from pathlib import Path
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def inline(s):
    s=html.escape(s.strip())
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',s)
    s=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',s)
    return s.replace('`','')


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--font',required=True)
    ap.add_argument('--input',default='docs/knowledge_mri/technical_report.md')
    a=ap.parse_args()
    src=Path(a.input); pdfmetrics.registerFont(TTFont('CJK',a.font))
    pdfmetrics.registerFontFamily('CJK',normal='CJK',bold='CJK',italic='CJK',boldItalic='CJK')
    body=ParagraphStyle('Body',fontName='CJK',fontSize=10,leading=16,spaceAfter=7,wordWrap='CJK',allowWidows=0,allowOrphans=0,textColor=colors.HexColor('#253244'))
    cell=ParagraphStyle('Cell',parent=body,fontSize=8.7,leading=13,spaceAfter=0)
    headings={n:ParagraphStyle('H'+str(n),parent=body,fontSize={1:24,2:15,3:12}[n],leading={1:34,2:23,3:19}[n],spaceBefore=14,spaceAfter=9,keepWithNext=True,textColor=colors.HexColor('#176d63')) for n in (1,2,3)}
    code=ParagraphStyle('Code',parent=body,fontSize=8.6,leading=14,leftIndent=8,backColor=colors.HexColor('#f1f5f6'),borderPadding=6)
    doc=Document(); normal=doc.styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(10.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    for st in ['Title','Heading 1','Heading 2','Heading 3']:
        doc.styles[st].font.name='宋体';doc.styles[st]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    story=[];lines=src.read_text().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line: continue
        if line.startswith('```'):
            block=[]
            while i<len(lines) and not lines[i].startswith('```'): block.append(lines[i]);i+=1
            i+=1;story.append(Paragraph('<br/>'.join(html.escape(x) for x in block),code));doc.add_paragraph('\n'.join(block));continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            rows=[[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.match(r'^\|[\s:|\-]+\|$',r)]
            data=[[Paragraph(inline(c),cell) for c in r] for r in rows]
            widths=[(481/len(rows[0]))]*len(rows[0])
            t=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e3f1ee')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7f9fa')]),('VALIGN',(0,0),(-1,-1),'TOP'),('BOX',(0,0),(-1,-1),.4,colors.HexColor('#c8d5db')),('INNERGRID',(0,0),(-1,-1),.25,colors.HexColor('#dbe2e7')),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            story.extend([t,Spacer(1,10)])
            dt=doc.add_table(rows=0,cols=len(rows[0]));dt.style='Light Shading Accent 1'
            for r in rows:
                for c,v in zip(dt.add_row().cells,r):c.text=v.replace('`','')
            continue
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if m:
            path=src.parent/m.group(2);im=Image(str(path));im.drawHeight=425*im.imageHeight/im.imageWidth;im.drawWidth=425
            story.append(im);story.append(Spacer(1,8));doc.add_picture(str(path),width=Inches(6.3));continue
        m=re.match(r'^(#{1,3}) (.*)',line)
        if m:
            n=len(m[1]);story.append(Paragraph(inline(m[2]),headings[n]));doc.add_heading(m[2],level=0 if n==1 else n-1);continue
        text=line
        story.append(Paragraph(inline(text),body));doc.add_paragraph(text.replace('`','').replace('**',''))
    def decorate(canvas,document):
        canvas.saveState();canvas.setStrokeColor(colors.HexColor('#dbe4e6'));canvas.line(57,804,538,804)
        canvas.setFont('CJK',8);canvas.setFillColor(colors.HexColor('#657484'))
        canvas.drawString(57,814,'WeKnora · 课题四 · 技术报告')
        canvas.drawString(57,30,'张凌浩 / zlh123123 · 2026-09-10')
        canvas.drawRightString(538,30,str(document.page));canvas.restoreState()
    output=src.with_suffix('.pdf')
    SimpleDocTemplate(str(output),pagesize=(595,842),leftMargin=57,rightMargin=57,topMargin=49,bottomMargin=53,title='WeKnora 知识网络与引导式学习 - 张凌浩',author='张凌浩 / zlh123123').build(story,onFirstPage=decorate,onLaterPages=decorate)
    doc.core_properties.author='张凌浩 / zlh123123';doc.core_properties.title='WeKnora 知识网络与引导式学习'
    doc.save(src.with_suffix('.docx'));print(output)
if __name__=='__main__':main()
